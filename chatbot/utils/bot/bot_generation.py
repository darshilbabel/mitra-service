import json
from concurrent.futures import ThreadPoolExecutor, as_completed

import pandas as pd
from celery import shared_task
from jinja2 import Template
from openpyxl import load_workbook
import logging
import os
import uuid
import re
from chatbot.llm_models.llm_script import handle_bedrock_model, handle_openai_model
from chatbot.models import CompanyBot, LLMProvider, CompanyBotTypeChoices, BotVernacular, Voice
from chatbot.services.core.prompt_builder import PromptBuilder
from chatbot.models import Company, CompanyStateMachine, OperationTypeChoices, LLMModel

logger = logging.getLogger('django')
AWS_KEY = os.getenv('AWS_ACCESS_KEY_ID')
AWS_SECRET_KEY = os.getenv('AWS_SECRET_ACCESS_KEY')

# =========================
# CONFIGURATION (EDIT HERE)
# =========================

MAX_WORKERS = 4
DEFAULT_BOT_NAME="Generated Bot"
DEFAULT_BOT_PROMPT="Sample Input prompt (If this is llm bot then please add bot persona)"

EXPECTED_COLUMNS = [
    "Q No.", "Section", "Main Question", "Expected responses",
    "Is Probing Required", "When to Probe", "Follow-up Questions",
    "What insight do you want from this question?"
]


def clean_string(value):
    if isinstance(value, str):
        return value.replace("\x00", "")
    return value


def normalize_flag(value):
    if not value:
        return ""
    value = str(value)
    value = value.replace("\xa0", "")  # remove NBSP
    value = re.sub(r'\s+', '', value)  # remove all whitespace
    return value.lower()


def convert_xlsx_to_json(file):
    import pandas as pd

    df = pd.read_excel(file)

    # Normalize column names
    df.columns = [str(col).strip() for col in df.columns]

    df["Q No."] = df["Q No."].ffill()
    df["Section"] = df["Section"].ffill()
    df["Main Question"] = df["Main Question"].ffill()

    # Detect instruction row safely using specific columns
    try:
        first_row = df.iloc[0]

        if (
            "category of the question" in str(first_row.get("Section", "")).lower()
            or "the actual question" in str(first_row.get("Main Question", "")).lower()
        ):
            df = df.iloc[1:]
    except Exception:
        pass

    # Reset index after possible removal
    df = df.reset_index(drop=True)

    # Forward fill
    df["Q No."] = df["Q No."].ffill()
    df["Section"] = df["Section"].ffill()
    df["Main Question"] = df["Main Question"].ffill()

    grouped = []

    for q_no, group in df.groupby("Q No."):
        base = group.iloc[0]

        item = {
            "Q No.": q_no,
            "Section": base["Section"],
            "Main Question": base["Main Question"],
            "Expected responses": base["Expected responses"],
            "What insight do you want from this question?": base["What insight do you want from this question?"],
            "probing": []
        }

        for _, row in group.iterrows():
            probing_flag = normalize_flag(row.get("Is Probing Required", ""))

            if probing_flag in {"true", "yes", "y", "1"}:
                item["probing"].append({
                    "When to Probe": row["When to Probe"],
                    "Follow-up Questions": row["Follow-up Questions"]
                })

        grouped.append(item)

    return grouped


def json_to_string(data, indent=0):
    space = "  " * indent
    output = ""

    try:
        if isinstance(data, dict):
            for key, value in data.items():
                output += f"{space}{key}:\n"
                output += json_to_string(value, indent + 1)

        elif isinstance(data, list):
            for i, item in enumerate(data):
                output += f"{space}- item {i + 1}:\n"
                output += json_to_string(item, indent + 1)

        else:
            output += f"{space}{data}\n"

    except Exception as e:
        output += f"{space}[ERROR SERIALIZING DATA: {e}]\n"

    return output


def validate_input(file):
    if not file:
        return {"error": "Please upload a file."}

    if not file.name.endswith(".xlsx"):
        return {"error": "Only .xlsx files are allowed."}

    return None


def validate_data(df):
    if "Main Question" not in df.columns:
        return {"error": "Missing Main Question column"}

    invalid_rows = df["Main Question"].isna() | (
        df["Main Question"].astype(str).str.strip() == ""
    )

    if invalid_rows.any():
        row_numbers = (df[invalid_rows].index + 2).tolist()
        return {
            "error": f"Empty Main Question at rows: {row_numbers}"
        }

    return None


def parse_excel(file):
    try:
        wb = load_workbook(file)
        ws = wb.active
    except Exception:
        return None, {"error": "Invalid or corrupted Excel file."}

    try:
        first_row = next(ws.iter_rows(min_row=1, max_row=1))
        columns = [cell.value for cell in first_row if cell.value]
        columns = [str(col).strip() for col in columns if col]
    except Exception:
        return None, {"error": "Unable to read columns from Excel file."}

    return columns, None


def validate_columns(uploaded_columns):
    missing = [col for col in EXPECTED_COLUMNS if col not in uploaded_columns]
    extra = [col for col in uploaded_columns if col not in EXPECTED_COLUMNS]

    if missing:
        return {"error": f"Missing columns: {', '.join(missing)}"}

    warning_msg = ""
    if extra:
        warning_msg = f"Extra columns will be ignored: {', '.join(extra)}"

    return {"warning": warning_msg}


def generate_bot_prompt(persona, file, opening_message=None, closing_message=None):
    # 1. Validate input
    error = validate_input(file)
    if error:
        return error

    # 2. Parse excel
    uploaded_columns, error = parse_excel(file)
    if error:
        return error

    # 3. Validate columns
    result = validate_columns(uploaded_columns)

    if "error" in result:
        return result

    file.seek(0)

    df = pd.read_excel(file)
    df.columns = [str(col).strip() for col in df.columns]
    df["Q No."] = df["Q No."].ffill()
    df["Section"] = df["Section"].ffill()
    df["Main Question"] = (
        df["Main Question"]
        .astype(str)
        .str.replace("\xa0", "")
        .str.strip()
        .replace("", pd.NA)
        .ffill()
    )

    data_error = validate_data(df)
    if data_error:
        return data_error

    file.seek(0)

    states_json = convert_xlsx_to_json(file)
    print("states_json: ", states_json)

    generate_question_prompts.delay(
        context_text=persona, states_json=states_json, opening_message=opening_message,
        closing_message=closing_message
    )

    warning_msg = result.get("warning", "")

    return {
        "success": True,
        "warning": warning_msg
    }


def handle_llm_model(company_bot, prompt, is_json_response=False):

    if company_bot.provider == LLMProvider.BEDROCK_CONVERSE:
        messages = [
            {
                'role': 'user',
                'content': [{'text': 'Generate the formatted prompt from the given input.'}]
            }
        ]
        response = handle_bedrock_model(
            system_prompt=prompt, messages=messages, model_name=company_bot.llm_model,
            temperature=company_bot.bot_temperature, max_token=company_bot.max_token, company_bot=company_bot,
            tools=None, top_p=company_bot.filter_score,
        )
    else:
        messages = [
            {
                "role": "user",
                "content": "Generate the formatted prompt from the given input."
            }
        ]
        response = handle_openai_model(
            system_prompt=prompt, messages=messages, model_name=company_bot.llm_model,
            temperature=company_bot.bot_temperature, max_token=company_bot.max_token,
            tool_choice='auto', is_json_response=is_json_response, company_bot=company_bot
        )

    return response


def generate_context(context_input):
    context_prompt = ""
    try:
        company_bot = CompanyBot.objects.filter(route='/bot_generation_context').first()
        if not company_bot:
            return {
                "bot_name": DEFAULT_BOT_NAME,
                "prompt": DEFAULT_BOT_PROMPT
            }
        prompt_builder = PromptBuilder()
        prompt_to_use = prompt_builder.build_system_prompt(
            company_bot=company_bot, other_data=f"Input: \n {context_input}"
        )
        response = handle_llm_model(company_bot=company_bot, prompt=prompt_to_use)

        bot_name = ""
        try:
            if isinstance(response, str):
                parsed = json.loads(response)
            else:
                parsed = response

            bot_name = parsed.get("bot_name", DEFAULT_BOT_NAME)
            context_prompt = parsed.get("prompt", DEFAULT_BOT_PROMPT)

        except Exception as e:
            print("Context JSON parsing failed:", e)

        return {"bot_name": bot_name, "prompt": context_prompt}
    except Exception as e:
        print(f"Context generation failed: {e}")
        return {
            "bot_name": DEFAULT_BOT_NAME,
            "prompt": DEFAULT_BOT_PROMPT
        }


def process_single_state(idx, state, global_context):
    try:
        state_str = json_to_string(state)

        company_bot = CompanyBot.objects.filter(route='/bot_generation_states').first()
        guard_rails_bot = CompanyBot.objects.filter(route='/bot_generation_guardrail').first()
        if not company_bot:
            print("No bot found for state generation")
            return {
                "index": idx,
                "name": state.get("Main Question", ""),
                "question": state.get("Main Question", ""),
                "section": state.get("Section", ""),
                "q_no": state.get("Q No."),
                "prompt": "",
                "status": "failed",
                "error": "Company bot not found"
            }

        raw_context = company_bot.context or ""
        context_data = {
            "guard_rails_context": getattr(guard_rails_bot, "context", ""),
        }

        try:
            rendered_context = Template(raw_context).render(context_data)
        except Exception as e:
            rendered_context = raw_context

        company_bot.context = rendered_context
        prompt_builder = PromptBuilder()
        other_data = {
            "Input Instruction": f"\n {state_str}",
            # "Global Context Already generated (no need to include these in state prompt again)": global_context,
            # "These are guard rails already included in the global prompt and no need to put them in state prompt "
            # "again or override them unless asked in [When to probe] condition.": f"{guard_rails_bot.context}"

        }

        prompt_to_use = prompt_builder.build_system_prompt(
            company_bot=company_bot, other_data=other_data
        )
        response=handle_llm_model(company_bot=company_bot, prompt=prompt_to_use)
        state_name = ""

        try:
            if isinstance(response, str):
                parsed_response = json.loads(response)
            else:
                parsed_response = response
            state_prompt = parsed_response.get("prompt", "") or parsed_response.get("response", "")
            state_name = parsed_response.get("state_name", "") or parsed_response.get("name", "")

        except Exception as e:
            print("JSON parsing failed:", e)
            print("Falling back to raw response")
            state_prompt = str(response)
        print("Response in process: ", response)
        return {
            "index": idx,
            "name": state_name or f"STATE_{idx + 1}",
            "question": state.get("Main Question", ""),
            "section": state.get("Section", ""),
            "q_no": state.get("Q No."),
            "prompt": state_prompt,
            "status": "success"
        }

    except Exception as e:
        print("Error Occurred processing single state: ", e)
        return {
            "index": idx,
            "name": state.get("Main Question", ""),
            "question": state.get("Main Question", ""),
            "section": state.get("Section", ""),
            "q_no": state.get("Q No."),
            "prompt": "",
            "status": "failed",
            "error": str(e)
        }


def generate_states_parallel(llm_states, global_context):
    results = []

    try:
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            futures = [
                executor.submit(process_single_state, idx, state, global_context)
                for idx, state in llm_states
            ]

            for future in as_completed(futures):
                try:
                    result = future.result()
                    results.append(result)
                except Exception as e:
                    print(f"Thread execution error: {e}")

    except Exception as e:
        print(f"Parallel execution failed: {e}")

    return results


def generate_unique_route_safe(bot_name):
    while True:
        route = generate_unique_route(bot_name)
        if not CompanyBot.objects.filter(route=route).exists():
            return route


def generate_unique_route(bot_name):

    slug = re.sub(r'[^a-zA-Z0-9]+', '-', bot_name.lower()).strip('-')
    unique_id = str(uuid.uuid4())[:8]

    return f"/{slug}-{unique_id}".replace("//", "/")


def save_bot_and_states(final_output, opening_message=None):
    try:
        print("Here is final output: ", final_output)
        company = Company.objects.get(id=1)
        main_bot_output = final_output.get("main_bot", {})
        unique_route = generate_unique_route_safe(
            bot_name=main_bot_output.get("bot_name", DEFAULT_BOT_NAME)
        )

        guard_rails_bot = CompanyBot.objects.filter(route='/bot_generation_guardrail').first()
        context_bot = CompanyBot.objects.filter(route='/bot_generation_context').first()
        tool_context, other_params, model_pricing = {}, {}, {}
        voices_data = []

        if context_bot and context_bot.other_params:
            tool_context = context_bot.other_params.get('tool_to_set', {})
            voices_data = context_bot.other_params.get('voice_to_set', {})
            model_pricing = context_bot.other_params.get('model_pricing', {})

        if model_pricing:
            other_params["model_pricing"] = model_pricing
        bot = CompanyBot.objects.create(
            name=main_bot_output.get("bot_name", DEFAULT_BOT_NAME).strip(),
            route=unique_route,
            max_token=8192,
            filter_score=0,
            connect_timeout=10,
            read_timeout=30,
            other_params=other_params,
            company=company,
            context=clean_string(main_bot_output.get("prompt", DEFAULT_BOT_PROMPT).strip()),
            llm_model=LLMModel.LLAMA_3_3_70B_INSTRUCT,
            bot_type=CompanyBotTypeChoices.STATE_MACHINE,
            pre_context=guard_rails_bot.context.strip() if guard_rails_bot and guard_rails_bot.context else None,
            tool_context=json.dumps(tool_context)
        )

        if bot and opening_message and opening_message.strip():
            BotVernacular.objects.create(
                company_bot=bot,
                language="en",
                name=bot.name,
                introductory_message=opening_message.strip()
            )

        if voices_data and isinstance(voices_data, list):
            for voice in voices_data:
                Voice.objects.create(
                    company_bot=bot,
                    type=voice.get("type"),
                    provider=voice.get("provider"),
                    name=voice.get("name"),
                    sample_link=voice.get("sample_link"),
                    language=voice.get("language"),
                    provider_code=voice.get("provider_code"),
                    gender=voice.get("gender"),
                    voice_speed=voice.get("voice_speed", 1.0),
                    other_params=voice.get("other_params")  # keep as is
                )

        state_info = final_output.get("state_info", [])

        step = 1

        for state in state_info:
            state_prompt = clean_string((state.get("prompt") or "").strip())
            state_name = (
                (state.get("name") or f"STATE_{step}").strip()
                if state_prompt
                else f"STATE_{step}"
            )
            state_name = clean_string(state_name)
            bot_question = clean_string((state.get("question") or "").strip())

            if state and isinstance(state, dict):
                CompanyStateMachine.objects.create(
                    company_bot=bot,
                    name=state_name,
                    bot_question=bot_question,
                    context=state_prompt,
                    step=step,
                    operation_type=(
                        OperationTypeChoices.LLM
                        if state.get("prompt")
                        else OperationTypeChoices.NON_LLM
                    )
                )
                step += 1

        print("Bot and states saved successfully")

        return {
            "status": True,
            "message": "Bot and states saved successfully"
        }

    except Exception as e:
        print(f"Error saving bot: {e}")
        try:
            CompanyBot.objects.filter(route='test-route-123').delete()
        except Exception as bot_delete_error:
            print(f"Error deleting bot: {e}")
        return {
            "status": False,
            "message": "Error saving bot"
        }


@shared_task
def generate_question_prompts(context_text, states_json, opening_message=None, closing_message=None):
    print("🚀 Starting prompt generation...\n")

    save_response = {
        "status": False,
        "message": "Prompt generation failed!!"
    }

    final_output = {
        "main_bot": {
            "bot_name": DEFAULT_BOT_NAME, "prompt": DEFAULT_BOT_PROMPT
        },
        "state_info": [],
        "metadata": {
            "status": "in_progress",
            "workers": MAX_WORKERS
        }
    }

    if context_text and context_text.strip():
        print("⚙️ Generating context...")
        context = generate_context(context_text)
        final_output["main_bot"] = context

        if not context:
            raise Exception("Context generation failed. Aborting pipeline.")

    if states_json:
        print("⚙️ Processing states...")

        llm_states = []
        non_llm_states = []

        for idx, state in enumerate(states_json):
            if state.get("probing"):
                llm_states.append((idx, state))
            else:
                non_llm_states.append({
                    "index": idx,
                    "name": state.get("Main Question", ""),
                    "question": state.get("Main Question", ""),
                    "section": state.get("Section", ""),
                    "q_no": state.get("Q No."),
                    "prompt": "",
                    "status": "skipped"
                })
        print("⚙️ Generating LLM states...")
        global_context = final_output["main_bot"].get("prompt", "")
        state_results = generate_states_parallel(llm_states, global_context)
        all_results = state_results + non_llm_states
        all_results.sort(key=lambda x: x["index"])

        for item in all_results:
            item.pop("index", None)

        final_output["state_info"] = all_results
        if opening_message and final_output["state_info"]:
            first_state = final_output["state_info"][0]

            existing_question = (first_state.get("question") or "").strip()
            opening_clean = opening_message.strip()

            if (opening_clean.lower() not in existing_question.lower() and
                    existing_question.lower() not in opening_clean.lower()
            ):
                updated_question = f"{opening_clean}\n{existing_question}".strip()
                first_state["question"] = updated_question

    final_output["metadata"]["status"] = "completed"
    closing_message = (closing_message or "").strip()

    if closing_message:
        final_output["state_info"].append({
            "name": "CLOSING",
            "question": closing_message,
            "prompt": "",
            "status": "static"
        })

    print("\n💾 Saving output...")

    save_response = save_bot_and_states(final_output, opening_message)

    return save_response
