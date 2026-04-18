import os

from django.contrib import admin
from simple_history.admin import SimpleHistoryAdmin
from .generic_upload_admin import BatchUploadMixin
from chatbot.filter.custom_date_from_filter import CustomAdvanceDateFilter
from chatbot.models import Profile, ProfileType, CompanyBot, CompanyBotTypeChoices, Voice
from chatbot.models.company_models import CompanyStateMachine
from django.shortcuts import redirect
from django.contrib import messages
from django.urls import path
from django.http import HttpResponseRedirect
from django.urls import reverse
from django.template.response import TemplateResponse
from operator import attrgetter
from chatbot.models import HistoricalCompanyStateMachine, HistoricalCompanyBot, HistoricalVoice
import difflib
from django.http import JsonResponse
from ..form.bot_form import CompanyBotAdminForm
from ..utils.bot.bot_generation import generate_bot_prompt


class CompanyStateMachineAdmin(admin.TabularInline):
    model = CompanyStateMachine
    fk_name = 'company_bot'
    extra = 1
    raw_id_fields = ['preprocess_bot', 'postprocess_bot']
    fields = (
        'name', 'step', 'use_stage_chats', 'text_conversion_type',
        'bot_question', 'completion_criteria', 'context', 'tool_context',
        'operation_type', 'skip_if_authenticated',
        'preprocess_type', 'preprocess_prompt', 'preprocess_bot', 'preprocess_output_mode',
        'postprocess_type', 'postprocess_prompt', 'postprocess_bot', 'postprocess_output_mode',
        'skip_to_step',
    )
    exclude = ('type',)

    def get_queryset(self, request):
        qs = super().get_queryset(request)
        return qs.order_by('step')


class VoiceProviderAdmin(admin.TabularInline):
    model = Voice
    extra = 1

    def get_queryset(self, request):
        qs = super().get_queryset(request)
        return qs.order_by('type', 'language')

    def formfield_for_dbfield(self, db_field, request, **kwargs):
        if db_field.name == "other_params":
            kwargs["help_text"] = "Leave empty to auto-load provider defaults."

        return super().formfield_for_dbfield(db_field, request, **kwargs)


@admin.register(CompanyBot)
class CompanyBotAdmin(BatchUploadMixin, SimpleHistoryAdmin):
    form = CompanyBotAdminForm
    list_display = ('name', 'company', 'created_at')
    list_filter = (
        'company',
        'name',
        'provider',
        'llm_model',
        CustomAdvanceDateFilter,
    )
    search_fields = ('name', 'company__name')
    date_hierarchy = 'created_at'
    ordering = ('-created_at',)
    inlines = [VoiceProviderAdmin]
    actions = ['duplicate_bot', 'export_selected_bots']

    enable_batch_upload = True
    batch_load_foreign_keys = True
    batch_upload_fields = ['name', 'company', 'provider', 'llm_model', 'context', 'max_token', 'route']

    import_template_name = 'admin/import_export/import.html'
    export_template_name = 'admin/import_export/export.html'

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                'export/',
                self.admin_site.admin_view(self.export_view),
                name='chatbot_companybot_export',
            ),
            path(
                'import/',
                self.admin_site.admin_view(self.import_view),
                name='chatbot_companybot_import',
            ),
            path(
                "<int:object_id>/revert/<str:model>/<int:history_id>/",
                self.admin_site.admin_view(self.revert_view),
                name="companybot_revert",
            ),
            path(
                "diff/<str:model>/<int:history_id>/",
                self.admin_site.admin_view(self.diff_view),
                name="companybot_diff",
            ),
            path(
                'generate-bot/',
                self.admin_site.admin_view(self.generate_bot_view),
                name='chatbot_companybot_generate_bot',
            ),
            path(
                'download-sample/',
                self.admin_site.admin_view(self.download_sample),
                name='chatbot_companybot_download_sample',
            ),
        ]
        return custom_urls + urls

    def diff_view(self, request, model, history_id):

        if model == "bot":
            record = HistoricalCompanyBot.objects.get(history_id=history_id)
            model_class = CompanyBot

        elif model == "voice":
            record = HistoricalVoice.objects.get(history_id=history_id)
            model_class = Voice

        elif model == "state":
            record = HistoricalCompanyStateMachine.objects.get(history_id=history_id)
            model_class = CompanyStateMachine

        else:
            return JsonResponse({"html": "<h3>Invalid model</h3>"})

        prev = model_class.history.filter(
            id=record.id,
            history_date__lt=record.history_date
        ).order_by("-history_date").first()

        diff_html = ""

        if prev:
            delta = record.diff_against(prev)

            for change in delta.changes:
                old = "" if change.old is None else str(change.old)
                new = "" if change.new is None else str(change.new)

                diff = difflib.HtmlDiff().make_table(
                    old.splitlines(),
                    new.splitlines(),
                    fromdesc="Old",
                    todesc="New",
                    context=True,
                    numlines=2
                )

                diff_html += f"<h3 style='margin-top:25px;'>Field: {change.field}</h3>{diff}"

        return JsonResponse({"html": diff_html})

    def revert_view(self, request, object_id, model, history_id):

        if model == "bot":
            history = HistoricalCompanyBot.objects.get(history_id=history_id)
            instance = CompanyBot.objects.get(pk=history.id)

        elif model == "voice":
            history = HistoricalVoice.objects.get(history_id=history_id)
            instance = Voice.objects.get(pk=history.id)

        elif model == "state":
            history = HistoricalCompanyStateMachine.objects.get(history_id=history_id)
            instance = CompanyStateMachine.objects.get(pk=history.id)

        else:
            self.message_user(request, "Invalid revert target.", level=messages.ERROR)
            return redirect(f"/admin/chatbot/companybot/{object_id}/change/")

        for field in history._meta.fields:

            if field.name in [
                "id", "history_id", "history_date", "history_user", "history_type",
            ]:
                continue

            if field.primary_key:
                continue

            setattr(instance, field.name, getattr(history, field.name))

        if model == "state":
            exists = CompanyStateMachine.objects.filter(
                company_bot=instance.company_bot,
                step=instance.step
            ).exclude(pk=instance.pk).exists()

            if exists:
                self.message_user(
                    request,
                    f"Cannot revert: step {instance.step} already exists.",
                    level=messages.ERROR
                )
                return redirect(f"/admin/chatbot/companybot/{object_id}/change/")

        instance.save()

        self.message_user(
            request,
            "Successfully reverted to selected version.",
            level=messages.SUCCESS
        )

        return redirect(f"/admin/chatbot/companybot/{object_id}/change/")

    def export_view(self, request):
        """Handle export requests"""
        from chatbot.views.admin.bot_admin_views import export_bots
        return export_bots(request)

    def import_view(self, request):
        """Handle import requests"""
        from chatbot.views.admin.bot_admin_views import import_bots
        return import_bots(request)

    def generate_bot_view(self, request):
        print("Here")
        print("request.method: ", request.method)
        if request.method == "POST":
            persona = request.POST.get("persona")
            opening_message = request.POST.get("opening_message")
            closing_message = request.POST.get("closing_message")
            file = request.FILES.get("questions_file")

            print("Opening message: ", opening_message)
            print("Closing message: ", closing_message)

            result = generate_bot_prompt(persona, file, opening_message, closing_message)

            return JsonResponse(result)

        context = dict(
            self.admin_site.each_context(request),
            title="Generate Bot",
        )

        return TemplateResponse(
            request,
            "admin/bot/generate_bot.html",
            context
        )

    def download_sample(self, request):
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.styles import Font
        from django.http import HttpResponse

        wb = Workbook()
        ws = wb.active
        ws.title = "Questions"

        # =========================
        # Row 1: Column Headers
        # =========================
        columns = [
            "Q No.", "Section", "Main Question", "Expected responses",
            "Is Probing Required", "When to Probe", "Follow-up Questions",
            "What insight do you want from this question?"
        ]
        ws.append(columns)

        # =========================
        # Row 2: Instructions
        # =========================
        ws.append([
            "#",
            "Category of the question (used for grouping, this is optional).",
            "The actual question the bot will ask the user (Mandatory).",
            "Typical or possible answers users might give.",
            "No",  # must match dropdown
            "Condition when follow-up should be asked (e.g., unclear or incomplete answer).",
            "Question to ask when probing is triggered.",
            "Purpose of asking this question (optional)."
        ])

        # =========================
        # Styling
        # =========================
        # Bold header
        for cell in ws[1]:
            cell.font = Font(bold=True)

        # Italic instruction row
        for cell in ws[2]:
            cell.font = Font(italic=True)

        # =========================
        # Dropdown (Column E)
        # =========================
        dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
        ws.add_data_validation(dv)

        for row in range(2, 200):
            dv.add(f"E{row}")

        # =========================
        # Freeze top rows
        # =========================
        ws.freeze_panes = "A3"

        # =========================
        # Response
        # =========================
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=sample_questions.xlsx'

        wb.save(response)
        return response

    def get_import_formats(self):
        """Define allowed import formats"""
        from import_export.formats import base_formats
        return [base_formats.CSV, base_formats.XLSX, base_formats.JSON]

    def get_export_formats(self):
        """Define allowed export formats"""
        from import_export.formats import base_formats
        return [base_formats.CSV, base_formats.XLSX, base_formats.JSON]

    def get_export_filename(self, request, queryset, file_format):
        """Generate filename for exports"""
        import datetime
        date_str = datetime.datetime.now().strftime('%Y-%m-%d')
        filename = f"company_bots_{date_str}"
        return f"{filename}.{file_format.get_extension()}"

    def get_queryset(self, request):
        qs = super().get_queryset(request)
        user_email = request.user.email
        profile = Profile.objects.filter(email=user_email)
        if request.user.is_superuser:
            return qs
        elif len(profile) > 0 and profile[0].profile_type == ProfileType.MODERATOR:
            return qs.filter(company=profile[0].company)
        else:
            return qs.none()

    def get_form(self, request, obj=None, **kwargs):
        form = super().get_form(request, obj, **kwargs)
        user = request.user
        user_email = request.user.email
        profile = Profile.objects.filter(email=user_email)
        if not user.is_superuser and len(profile) > 0 and profile[0].profile_type == ProfileType.MODERATOR:
            company_field = form.base_fields.get('company')
            if company_field:
                form.base_fields['company'].queryset = form.base_fields['company'].queryset.filter(
                    id=profile[0].company.id)
            form.base_fields = {field_name: form.base_fields[field_name] for field_name in form.base_fields}

        if form.base_fields.get('context'):
            form.base_fields['context'].label = 'Prompt'
        if form.base_fields.get('pre_context'):
            form.base_fields['pre_context'].label = 'Guardrails'

        form.base_fields = {field_name: form.base_fields[field_name] for field_name in form.base_fields}
        return form

    def changeform_view(self, request, object_id=None, form_url='', extra_context=None):
        if object_id:
            obj = self.model.objects.get(pk=object_id)
            if obj.bot_type == CompanyBotTypeChoices.STATE_MACHINE:
                self.inlines = [VoiceProviderAdmin, CompanyStateMachineAdmin]

            else:
                self.inlines = [VoiceProviderAdmin]
        else:
            self.inlines = [VoiceProviderAdmin]
        return super().changeform_view(request, object_id, form_url, extra_context)

    def history_view(self, request, object_id, extra_context=None):
        obj = self.get_object(request, object_id)

        bot_history = list(obj.history.all())
        sm_history = list(CompanyStateMachine.history.filter(company_bot=obj))
        voice_history = list(Voice.history.filter(company_bot=obj))

        bot_history.sort(key=attrgetter("history_date"), reverse=True)
        sm_history.sort(key=attrgetter("history_date"), reverse=True)
        voice_history.sort(key=attrgetter("history_date"), reverse=True)

        history = bot_history + sm_history + voice_history

        latest_bot_marked = False
        latest_steps = {}

        import difflib

        for record in history:
            try:
                model_name = record._meta.model_name

                if model_name == "historicalcompanybot":
                    if not latest_bot_marked:
                        record.is_latest = True
                        latest_bot_marked = True
                    else:
                        record.is_latest = False

                elif model_name == "historicalcompanystatemachine":
                    step = record.step

                    if step not in latest_steps:
                        record.is_latest = True
                        latest_steps[step] = True
                    else:
                        record.is_latest = False

                elif model_name == "historicalvoice":
                    if not hasattr(self, "_latest_voice"):
                        record.is_latest = True
                        self._latest_voice = True
                    else:
                        record.is_latest = False
                else:
                    record.is_latest = False

                record.changes = []
                record.diff_html = []

            except Exception:
                record.changes = []
                record.diff_html = []

        context = {
            **self.admin_site.each_context(request),
            "title": f"History: {obj}",
            "bot_history": bot_history,
            "sm_history": sm_history,
            "voice_history": voice_history,
            "object": obj,
            "opts": self.model._meta,
        }

        return TemplateResponse(
            request,
            "admin/combined_history.html",
            context,
        )

    def duplicate_bot(self, request, queryset):
        if queryset.count() != 1:
            self.message_user(request, "Please select exactly one bot to duplicate.", level=messages.ERROR)
            return

        original = queryset.first()

        new_bot = CompanyBot.objects.get(pk=original.pk)
        new_bot.pk = None
        new_bot.name = f"{original.name} (Copy)"
        new_bot.save()

        original_voice_providers = Voice.objects.filter(company_bot=original)
        for voice in original_voice_providers:
            voice.pk = None
            voice.company_bot = new_bot
            voice.save()

        if original.bot_type == CompanyBotTypeChoices.STATE_MACHINE:
            original_state_machines = CompanyStateMachine.objects.filter(company_bot=original)
            for sm in original_state_machines:
                sm.pk = None
                sm.company_bot = new_bot
                sm.save()

        self.message_user(request, "Bot duplicated successfully!", level=messages.SUCCESS)
        return redirect(f"/admin/chatbot/companybot/{new_bot.id}/change/")

    def export_selected_bots(self, request, queryset):
        """Custom export action"""
        selected_ids = queryset.values_list('id', flat=True)
        ids_str = ','.join(str(id) for id in selected_ids)

        info = self.model._meta.app_label, self.model._meta.model_name
        url = reverse('admin:%s_%s_export' % info) + f'?ids={ids_str}'
        return HttpResponseRedirect(url)

    export_selected_bots.short_description = "Export selected bots"

    def changelist_view(self, request, extra_context=None):
        """Add custom buttons to the changelist view"""
        extra_context = extra_context or {}
        extra_context['custom_buttons'] = True
        extra_context['show_generate_bot'] = os.getenv("SHOW_GENERATE_BOT", "False") == "True"
        return super().changelist_view(request, extra_context=extra_context)

    duplicate_bot.short_description = "Duplicate selected bot"


class HistoricalCompanyBotAdmin(admin.ModelAdmin):
    readonly_fields = [field.name for field in HistoricalCompanyBot._meta.fields]

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def has_change_permission(self, request, obj=None):
        return False


class HistoricalCompanyStateMachineAdmin(admin.ModelAdmin):
    readonly_fields = [field.name for field in HistoricalCompanyStateMachine._meta.fields]

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def has_change_permission(self, request, obj=None):
        return False


admin.site.register(HistoricalCompanyBot, HistoricalCompanyBotAdmin)
admin.site.register(HistoricalCompanyStateMachine, HistoricalCompanyStateMachineAdmin)
